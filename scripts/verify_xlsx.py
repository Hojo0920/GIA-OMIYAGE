#!/usr/bin/env python3
"""xlsx の数式を LibreOffice で再計算し、ファイルに書いたキャッシュ値と全数比較する。

使い方: python3 scripts/verify_xlsx.py deliverables/anto_omiyage_candidates.xlsx

LibreOffice は既定では xlsx の数式を再計算しない（キャッシュ値をそのまま表示する）ため、専用のプロファイルを作り、
「読み込み時に常に再計算」（OOXMLRecalcMode=0）にして変換する。普段使いの設定は変えない。
"""
import pathlib
import subprocess
import sys
import tempfile

import openpyxl

XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def main(path: str) -> int:
    src = pathlib.Path(path).resolve()
    with tempfile.TemporaryDirectory() as tmp:
        prof = pathlib.Path(tmp) / "profile" / "user"
        prof.mkdir(parents=True)
        (prof / "registrymodifications.xcu").write_text(XCU, encoding="utf-8")
        out = pathlib.Path(tmp) / "out"
        subprocess.run(["soffice", f"-env:UserInstallation=file://{prof.parent}", "--headless", "--convert-to", "xlsx",
                        "--outdir", str(out), str(src)], check=True, capture_output=True, timeout=300)
        recalc = out / src.name
        wf = openpyxl.load_workbook(src)
        wo = openpyxl.load_workbook(src, data_only=True)
        wr = openpyxl.load_workbook(recalc, data_only=True)
        total = bad = 0
        for ws in wf.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        total += 1
                        a = wo[ws.title][c.coordinate].value
                        b = wr[ws.title][c.coordinate].value
                        a = "" if a is None else a
                        b = "" if b is None else b
                        same = abs(a - b) < 1e-9 if isinstance(a, (int, float)) and isinstance(b, (int, float)) else str(a) == str(b)
                        if not same:
                            bad += 1
                            if bad <= 20:
                                print(f"不一致 {ws.title}!{c.coordinate} キャッシュ={a!r} 再計算={b!r} {c.value[:80]}")
        print(f"数式セル {total} 件、不一致 {bad} 件")
        return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
